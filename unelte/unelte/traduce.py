"""traduce – traducere prin AI cu păstrarea formatării.

Markdown: blocurile de cod, formulele și antetul YAML nu se trimit; codul și formulele inline
sunt înlocuite cu marcaje care trebuie să revină neschimbate. Cererea și răspunsul sunt JSON
cu schemă (listă de {id, text}); la o nepotrivire, bucata se retrimite o dată, apoi comanda
se oprește fără fișier parțial.
"""
from __future__ import annotations

import json
import re
import sys
import tomllib
from dataclasses import dataclass
from pathlib import Path

from . import ai
from .cli import parser
from .comun import Progres, config, fisiere, prompturi

CHUNK_CHARS = 6000
LANGS = {"ro": "română", "en": "engleză", "de": "germană", "fr": "franceză", "es": "spaniolă", "it": "italiană",
         "hu": "maghiară", "pt": "portugheză", "nl": "neerlandeză", "pl": "poloneză", "ru": "rusă", "uk": "ucraineană"}
SCHEMA = {
    "type": "object",
    "required": ["blocuri"],
    "properties": {"blocuri": {"type": "array", "items": {
        "type": "object", "required": ["id", "text"],
        "properties": {"id": {"type": "integer"}, "text": {"type": "string"}}}}},
}


class Mismatch(Exception):
    pass


@dataclass
class Segment:
    text: str
    translate: bool


# ---------- împărțirea Markdown ----------

_FRONT = re.compile(r"\A---[ \t]*\n.*?\n(?:---|\.\.\.)[ \t]*\n?", re.S)
_FENCE_OPEN = re.compile(r"^[ \t]*(```+|~~~+)")
_KEEP_BLOCK = re.compile(r"^\s*(\$\$.*\$\$|<!--.*-->|(?:[-*_][ \t]*){3,}|\[[^\]]+\]:\s*\S+.*)\s*$", re.S)
_PROTECT = re.compile(r"(`+)(?!`).+?(?<!`)\1(?!`)|(?<![\\$])\$(?!\s)[^$\n]+?(?<!\s)\$(?!\d)|\[@[^\]]+\]|<[^>\s]+>")


def split_markdown(text: str) -> list[Segment]:
    segs: list[Segment] = []
    m = _FRONT.match(text)
    pos = 0
    if m:
        segs.append(Segment(m.group(0), False))
        pos = m.end()
    lines = text[pos:].splitlines(keepends=True)
    i = 0
    block: list[str] = []

    def flush():
        if block:
            t = "".join(block)
            core = t.rstrip("\n")
            tail = t[len(core):]
            indented = all(ln.startswith(("    ", "\t")) or not ln.strip() for ln in block)
            keep = bool(_KEEP_BLOCK.match(core)) or indented
            segs.append(Segment(core, not keep and bool(core.strip())))
            if tail:
                segs.append(Segment(tail, False))
            block.clear()

    while i < len(lines):
        ln = lines[i]
        fence = _FENCE_OPEN.match(ln)
        if fence:
            flush()
            marker = fence.group(1)
            j = i + 1
            while j < len(lines) and not lines[j].lstrip().startswith(marker):
                j += 1
            segs.append(Segment("".join(lines[i:j + 1]), False))
            i = j + 1
            continue
        if ln.strip().startswith("$$") and not (ln.strip().endswith("$$") and len(ln.strip()) > 2):
            flush()
            j = i + 1
            while j < len(lines) and "$$" not in lines[j]:
                j += 1
            segs.append(Segment("".join(lines[i:j + 1]), False))
            i = j + 1
            continue
        if not ln.strip():
            flush()
            segs.append(Segment(ln, False))
        else:
            block.append(ln)
        i += 1
    flush()
    assert "".join(s.text for s in segs) == text
    return segs


def protect(text: str) -> tuple[str, list[str]]:
    kept: list[str] = []

    def sub(m):
        kept.append(m.group(0))
        return f"⟦{len(kept) - 1}⟧"
    return _PROTECT.sub(sub, text), kept


def restore(text: str, kept: list[str]) -> str:
    found = re.findall(r"⟦(\d+)⟧", text)
    if sorted(int(x) for x in found) != list(range(len(kept))):
        raise Mismatch("formulele sau codul inline nu au fost păstrate")
    return re.sub(r"⟦(\d+)⟧", lambda m: kept[int(m.group(1))], text)


# ---------- glosar ----------

def glossary(target: str) -> dict[str, str]:
    p = config.config_dir() / "glosar.toml"
    try:
        with p.open("rb") as f:
            data = tomllib.load(f)
    except (OSError, tomllib.TOMLDecodeError):
        return {}
    section = data.get(target) or {}
    return {str(k): str(v) for k, v in section.items() if isinstance(v, str)}


# ---------- traducere ----------

def group(units: list[tuple[int, str]], limit: int) -> list[list[tuple[int, str]]]:
    out, cur, size = [], [], 0
    for u in units:
        if cur and size + len(u[1]) > limit:
            out.append(cur)
            cur, size = [], 0
        cur.append(u)
        size += len(u[1])
    if cur:
        out.append(cur)
    return out


def translate_units(texts: list[str], target: str, source: str | None, model: str | None,
                    progress=None) -> list[str]:
    """Traduce o listă de texte, păstrând ordinea și numărul lor."""
    protected = [protect(t) for t in texts]
    units = [(i, p[0]) for i, p in enumerate(protected)]
    terms = glossary(target)
    gl = "\n".join(f"- „{k}” → „{v}”" for k, v in terms.items()) or "(glosar gol)"
    limit = ai.max_chunk_chars("traducere", CHUNK_CHARS, model)
    result: dict[int, str] = {}
    batches = group(units, limit)
    for k, batch in enumerate(batches, 1):
        if progress and len(batches) > 1:
            progress(f"Traduc bucata {k} din {len(batches)}")
        payload = json.dumps({"blocuri": [{"id": i, "text": t} for i, t in batch]}, ensure_ascii=False, indent=1)
        msgs = prompturi.mesaje("traducere", din_limba=f"din {LANGS.get(source, source)}" if source else
                                "din limba lor (detectată)", in_limba=LANGS.get(target, target), glosar=gl,
                                blocuri=payload)
        for attempt in (1, 2):
            r = ai.ruleaza("traducere", msgs, model=model, schema=SCHEMA, la_progres=progress)
            got = r.date["blocuri"]
            try:
                ids = [int(b["id"]) for b in got]
                if ids != [i for i, _t in batch]:
                    raise Mismatch(f"am trimis {len(batch)} blocuri și am primit {len(got)} (sau în altă ordine)")
                restored = {int(b["id"]): restore(str(b["text"]), protected[int(b["id"])][1]) for b in got}
            except Mismatch as e:
                if attempt == 2:
                    raise config.EroareUtilizare(f"Traducerea bucății {k} nu corespunde cu originalul ({e}). "
                                                 "Nu am scris niciun fișier.") from e
                continue
            result.update(restored)
            break
    return [result[i] for i in range(len(texts))]


def translate_markdown(text: str, target: str, source: str | None, model: str | None, progress=None) -> str:
    segs = split_markdown(text)
    todo = [s.text for s in segs if s.translate]
    done = iter(translate_units(todo, target, source, model, progress) if todo else [])
    return "".join(next(done) if s.translate else s.text for s in segs)


def translate_docx(src: Path, dest: Path, target: str, source: str | None, model: str | None, progress=None) -> int:
    from docx import Document
    doc = Document(str(src))
    paras = [p for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                paras += [p for p in cell.paragraphs if p.text.strip()]
    seen, unique = set(), []
    for p in paras:
        if id(p._p) not in seen:
            seen.add(id(p._p))
            unique.append(p)
    translated = translate_units([p.text for p in unique], target, source, model, progress)
    for p, t in zip(unique, translated):
        runs = p.runs
        if runs:
            runs[0].text = t
            for r in runs[1:]:
                r._r.getparent().remove(r._r)
        else:
            p.add_run(t)
    doc.save(str(dest))
    return len(unique)


def main(argv=None) -> int:
    ap = parser("traduce", "Traduce un document prin AI, păstrând formatarea Markdown sau stilurile Word.",
                "Exemple:\n  traduce articol.md --in en            # → articol.en.md\n"
                "  traduce rezumat.md --in ro --din en\n  traduce document.docx --in en         # → document.en.docx\n"
                "  echo \"text\" | traduce - --in en\n\n"
                "Glosarul (~/.config/unelte/glosar.toml) are câte o secțiune per limbă țintă:\n"
                "  [en]\n  \"active circulante\" = \"current assets\"")
    ap.add_argument("fisier", help="fișierul .md, .txt sau .docx, ori - pentru intrarea standard")
    ap.add_argument("--in", dest="target", required=True, metavar="LIMBA", help="limba țintă (en, ro, de, fr…)")
    ap.add_argument("--din", dest="source", metavar="LIMBA", help="limba sursă (implicit detectată)")
    ap.add_argument("--model", help="alias sau furnizor:model")
    a = ap.parse_args(argv)
    if a.model:
        ai.check_model_ref(a.model)
    progres = Progres()
    try:
        if a.fisier == "-":
            print(translate_markdown(sys.stdin.read(), a.target, a.source, a.model, progres), end="")
            return 0
        path = Path(a.fisier)
        if not path.is_file():
            raise config.EroareUtilizare(f"Fișierul {path} nu există.")
        if path.suffix.lower() == ".docx":
            print("Notă: în .docx se păstrează stilul fiecărui paragraf, dar formatarea din interiorul lui "
                  "(cuvinte îngroșate izolate, linkuri) se pierde.", file=sys.stderr)
            dest = fisiere.rezultat(path, a.target)
            tmp = dest.with_name(f".{dest.name}.partial")
            n = translate_docx(path, tmp, a.target, a.source, a.model, progres)
            tmp.replace(dest)
            progres.gata()
            print(f"Scris: {dest} ({n} paragrafe)")
            return 0
        text = path.read_text(encoding="utf-8")
        out = translate_markdown(text, a.target, a.source, a.model, progres)
        dest = fisiere.rezultat(path, a.target)
        fisiere.scrie_atomic(dest, out)
        progres.gata()
        print(f"Scris: {dest}")
        return 0
    finally:
        progres.gata()
