import subprocess
import zlib
import struct

import pytest
from conftest import needs, set_config
from docx import Document
from docx.shared import Cm, Pt

from unelte import converteste, verifica_format
from unelte.comun import config, profil, sistem


def png(path, w=40, h=30):
    raw = b"".join(b"\x00" + b"\x80\x40\x20" * w for _ in range(h))

    def ch(t, d):
        return struct.pack(">I", len(d)) + t + d + struct.pack(">I", zlib.crc32(t + d) & 0xffffffff)
    path.write_bytes(b"\x89PNG\r\n\x1a\n" + ch(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0))
                     + ch(b"IDAT", zlib.compress(raw)) + ch(b"IEND", b""))
    return path


def sty(doc, name):
    return next(s for s in doc.styles if s.name == name)


@pytest.fixture
def prof(env):
    return profil.incarca("facultate")


def good_doc(prof, tmp_path, body_paragraphs=3):
    """Un document corect, pornind de la reference.docx al profilului."""
    doc = Document(str(prof.reference))
    for p in list(doc.paragraphs):
        p._p.getparent().remove(p._p)
    doc.add_paragraph("Introducere", style=sty(doc, "Heading 1"))
    for i in range(body_paragraphs):
        doc.add_paragraph(f"Paragraful {i} al lucrării, cu text suficient pentru verificare.", style=sty(doc, "Body Text"))
    doc.add_paragraph("Detalii", style=sty(doc, "Heading 2"))
    doc.add_picture(str(png(tmp_path / "f.png")))
    doc.add_paragraph("Figura 1. Evoluția indicatorului", style=sty(doc, "Image Caption"))
    doc.add_paragraph("Bibliografie", style=sty(doc, "Heading 1"))
    doc.add_paragraph("Popescu, I. (2020). Analiză financiară.", style=sty(doc, "Bibliography"))
    return doc


def test_document_corect(prof, tmp_path):
    doc = good_doc(prof, tmp_path)
    p = tmp_path / "ok.docx"
    doc.save(str(p))
    rep = verifica_format.check_docx(p, prof)
    assert rep.problems == [], rep.text()


def test_probleme_cunoscute(prof, tmp_path):
    doc = good_doc(prof, tmp_path, body_paragraphs=0)
    # font moștenit din stil: stilul are Arial 11, run-urile nu au font propriu
    st = doc.styles.add_style("Corp greșit", 1)
    st.base_style = next(s for s in doc.styles if s.name == "Body Text")
    st.font.name = "Arial"
    st.font.size = Pt(11)
    body = doc.paragraphs[0]._p
    for i in range(140):
        par = doc.add_paragraph(f"Paragraf cu font greșit numărul {i}.", style=sty(doc, "Corp greșit"))
        body.addnext(par._p)
        body = par._p
    doc.add_paragraph("Sub-sub-titlu sărit", style=sty(doc, "Heading 3"))          # după Heading 1 (Bibliografie)
    doc.add_picture(str(png(tmp_path / "g.png")))                       # imagine fără legendă, la final
    doc.sections[0].left_margin = Cm(2.0)
    p = tmp_path / "rau.docx"
    doc.save(str(p))
    rep = verifica_format.check_docx(p, prof)
    text = rep.text()
    assert "marginea stânga 2.00 cm; cerut: 3.00 cm" in text
    assert "140 de paragrafe au Arial 11 în loc de Times New Roman 12" in text
    assert "(primul: §2 «Paragraf cu font greșit numărul 0.»)" in text
    assert "Heading 3 după Heading 1" in text
    assert "imagine fără legendă" in text
    assert text.count("Arial 11") == 1                                   # problemele repetate sunt comprimate


def test_rezolvare_prin_lantul_de_stiluri(prof, tmp_path):
    doc = Document(str(prof.reference))
    for p in list(doc.paragraphs):
        p._p.getparent().remove(p._p)
    base = doc.styles.add_style("Baza", 1)
    base.font.name = "Georgia"
    base.font.size = Pt(13)
    child = doc.styles.add_style("Copil", 1)
    child.base_style = base                       # Copil nu are font: îl moștenește din Baza
    child.paragraph_format.line_spacing = 2.0
    para = doc.add_paragraph("Text moștenit", style=sty(doc, "Copil"))
    run_p = doc.add_paragraph("", style=sty(doc, "Body Text"))
    run_p.add_run("Run cu font propriu").font.name = "Courier New"
    plain = doc.add_paragraph("Doar docDefaults")
    p = tmp_path / "lant.docx"
    doc.save(str(p))
    d = Document(str(p))
    styles = verifica_format.Styles(d)
    e = verifica_format.effective(d.paragraphs[0], styles)
    assert (e.font, e.size, e.line_spacing) == ("Georgia", 13, 2.0)
    assert verifica_format.effective(d.paragraphs[1], styles).font == "Courier New"
    e3 = verifica_format.effective(d.paragraphs[2], styles)
    assert e3.font == "Times New Roman" and e3.size == 12            # din Normal / docDefaults ale profilului
    del para, run_p, plain


def test_markdown_citari_si_antet(env, prof, tmp_path):
    bib = tmp_path / "b.bib"
    bib.write_text("@book{pop2020, title={A}}\n@comment{ignorat,}\n@article{ion:2021, title={B}}\n")
    set_config(env, f'[documente]\nbibliografie = "{bib}"\n')
    md = tmp_path / "l.md"
    md.write_text("---\ntitlu: T\nautor: A\n---\n\nText [@pop2020; @lipsa, p. 3] și @ion:2021.\n"
                  "Email ana@exemplu.ro nu e citare. `@cod` nici.\n")
    rep = verifica_format.check_markdown(md, prof, config.load())
    lines = rep.text()
    assert "@lipsa: cheie inexistentă" in lines and "@pop2020" not in lines.split("Info")[0]
    assert "lipsește câmpul «coordonator»" in lines
    assert "Chei lipsă (1): lipsa" in lines
    assert verifica_format.citation_keys("[@a, p. 1] [-@b] @c.") == ["a", "b", "c"]


def test_argumentele_pandoc_per_profil(env, prof, tmp_path, monkeypatch):
    calls = []

    def fake(args, **kw):
        calls.append(args)
        if args[0] == "pandoc":
            Document(str(prof.reference)).save(args[args.index("--output") + 1])
        return subprocess.CompletedProcess(args, 0, stdout="", stderr="")
    monkeypatch.setattr(sistem, "ruleaza", fake)
    monkeypatch.setattr(sistem, "necesita", lambda p: p)
    bib = tmp_path / "b.bib"
    bib.write_text("@book{pop2020, title={A}}\n")
    set_config(env, f'[documente]\nbibliografie = "{bib}"\n')
    md = tmp_path / "lucrare.md"
    md.write_text("---\ntitlu: T\ncuprins: true\n---\n\n# Unu\n\nText [@pop2020] și [@absent].\n")
    out = []
    dest, _pdf, _rep = converteste.convert(md, prof, config.load(), out=out.append)
    args = calls[0]
    assert dest == tmp_path / "lucrare.docx"
    assert args[args.index("--reference-doc") + 1] == str(prof.reference)
    assert args[args.index("--csl") + 1] == str(prof.csl)
    assert "--citeproc" in args and args[args.index("--bibliography") + 1] == str(bib)
    assert "--lua-filter" in args and "--toc" not in args              # cuprinsul îl pune filtrul
    assert "--number-sections" in args
    assert "lang=ro-RO" in args and "reference-section-title=Bibliografie" in args
    assert args[args.index("--resource-path") + 1] == str(tmp_path.resolve())
    assert any("citări fără intrare" in o and "absent" in o for o in out)
    calls.clear()
    converteste.convert(md, prof, config.load(), no_title=True, verify=False, out=out.append)
    assert "--toc" in calls[0] and "--lua-filter" not in calls[0]
    # fișier deschis în Word: rezultatul merge alături
    (tmp_path / "~$lucrare.docx").write_text("")
    assert converteste.output_path(md, ".docx").name == "lucrare (2).docx"


@needs("pandoc")
def test_conversie_reala_trece_verificarea(env, prof, tmp_path):
    png(tmp_path / "fig.png")
    bib = tmp_path / "b.bib"
    bib.write_text("@book{pop2020, author={Popescu, Ion}, title={Analiză}, publisher={Ed}, year={2020}}\n")
    set_config(env, f'[documente]\nbibliografie = "{bib}"\n')
    md = tmp_path / "l.md"
    md.write_text("---\nuniversitate: U\nfacultate: F\ntitlu: Titlu\nautor: A\ncoordonator: C\noras: O\nan: 2026\n---\n\n"
                  "# Introducere\n\nText cu formulă $x^2$ și citare [@pop2020].\n\n$$a = b$$\n\n"
                  "![Figura de test](fig.png)\n\n## Secțiune\n\n- unu\n- doi\n")
    dest, _p, rep = converteste.convert(md, prof, config.load(), out=lambda s: None)
    assert rep.problems == [], rep.text()
    text = "\n".join(p.text for p in Document(str(dest)).paragraphs)
    assert "Popescu" in text and "Coordonator științific: C" in text
