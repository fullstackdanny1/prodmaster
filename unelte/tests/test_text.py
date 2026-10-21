"""corecteaza (cu LanguageTool fals) și traduce (cu AI fals)."""
import json
import threading
import urllib.parse
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

import pytest
from conftest import set_config

from unelte import corecteaza, traduce

MD = """---
titlu: Test greșit
---

# Titlu

Acesta e un text cu o greseala. Formula $a + gresit = b$ nu se verifică.

```python
greseala_in_cod = 1
```

Vezi [linkul greșit](https://exemplu.ro/greseala) și `cod greseala`.
"""


class FakeLT:
    """LanguageTool fals: semnalează fiecare apariție a unui cuvânt din `bad`."""

    def __init__(self, bad=("greseala",)):
        self.bad = bad
        self.requests = []
        server = self

        class H(BaseHTTPRequestHandler):
            def log_message(self, *a):
                pass

            def do_POST(self):
                n = int(self.headers.get("Content-Length", 0))
                form = urllib.parse.parse_qs(self.rfile.read(n).decode())
                text = form["text"][0]
                server.requests.append({k: v[0] for k, v in form.items()})
                matches = []
                for w in server.bad:
                    start = 0
                    while (i := text.find(w, start)) >= 0:
                        matches.append({"message": "Posibilă greșeală de ortografie", "offset": i, "length": len(w),
                                        "replacements": [{"value": "greșeală"}],
                                        "rule": {"id": "MORFOLOGIK_RULE_RO_RO", "issueType": "misspelling"}})
                        start = i + 1
                body = json.dumps({"matches": matches}).encode()
                self.send_response(200)
                self.send_header("Content-Type", "application/json")
                self.end_headers()
                self.wfile.write(body)

        self.httpd = ThreadingHTTPServer(("127.0.0.1", 0), H)
        threading.Thread(target=self.httpd.serve_forever, args=(0.05,), daemon=True).start()

    @property
    def url(self):
        return f"http://127.0.0.1:{self.httpd.server_address[1]}"

    def stop(self):
        self.httpd.shutdown()


@pytest.fixture
def lt(env):
    s = FakeLT()
    set_config(env, f'[corecteaza]\nserver = "{s.url}"\n')
    yield s
    s.stop()


def test_mascarea_pastreaza_pozitiile():
    masked = corecteaza.mask_markdown(MD)
    assert len(masked) == len(MD)
    assert masked.count("\n") == MD.count("\n")
    assert "greseala" not in masked.split("# Titlu")[0]                # antetul
    assert "gresit" not in masked                                      # formula, URL-ul
    assert masked.count("greseala") == 1                               # doar cea din text
    i = masked.index("greseala")
    assert MD[i:i + 8] == "greseala"
    assert "linkul greșit" in masked and "https" not in masked


def test_raport_doar_text_natural(lt, tmp_path):
    p = tmp_path / "t.md"
    p.write_text(MD)
    assert corecteaza.main([str(p)]) == 1
    from unelte.comun import config
    found = corecteaza.check_text(MD, config.load(), "ro-RO", set(), set())
    assert len(found) == 1 and MD[found[0].offset:found[0].offset + 8] == "greseala"
    assert lt.requests[0]["language"] == "ro-RO"


def test_dictionar_personal(lt, tmp_path, capsys):
    assert corecteaza.main(["--adauga-cuvant", "greseala"]) == 0
    p = tmp_path / "t.md"
    p.write_text("Un text cu greseala.\n")
    assert corecteaza.main([str(p)]) == 0
    assert "Nicio greșeală" in capsys.readouterr().out


def test_bucati_lungi_cu_pozitii_ajustate(lt):
    para = ("Propoziție corectă fără probleme. " * 30).strip()
    text = "\n\n".join([para] * 40) + "\n\nIar aici o greseala.\n"
    from unelte.comun import config
    found = corecteaza.check_text(text, config.load(), "ro-RO", set(), set(), masked=False)
    assert len(lt.requests) > 1 and all(len(r["text"]) <= corecteaza.CHUNK_LIMIT for r in lt.requests)
    assert len(found) == 1 and text[found[0].offset:found[0].offset + 8] == "greseala"
    pieces = corecteaza.chunks(text, 5000)
    assert "".join(p for _s, p in pieces) == text and all(text[s:s + len(p)] == p for s, p in pieces)


def test_interactiv_nu_suprascrie(lt, tmp_path, monkeypatch):
    p = tmp_path / "t.md"
    p.write_text("O greseala și altă greseala.\n")
    answers = iter(["1", "s"])
    monkeypatch.setattr("builtins.input", lambda *_: next(answers))
    corecteaza.main([str(p), "--interactiv"])
    assert p.read_text() == "O greseala și altă greseala.\n"
    assert (tmp_path / "t.corectat.md").read_text() == "O greșeală și altă greseala.\n"
    answers2 = iter(["1", "1"])
    monkeypatch.setattr("builtins.input", lambda *_: next(answers2))
    corecteaza.main([str(p), "--interactiv"])
    assert (tmp_path / "t.corectat (2).md").read_text() == "O greșeală și altă greșeală.\n"


def test_stil_prin_ai(lt, fake_ai, tmp_path, capsys):
    p = tmp_path / "t.md"
    p.write_text("Textul este scris un pic cam greoi.\n")
    lt.bad = ()
    fake_ai.reply("Textul e scris cam greoi.")
    corecteaza.main([str(p), "--stil"])
    assert (tmp_path / "t.stil.md").read_text() == "Textul e scris cam greoi.\n"
    assert "- este  →  e" in capsys.readouterr().out


def test_server_oprit(env, tmp_path):
    set_config(env, '[corecteaza]\nserver = "http://127.0.0.1:9"\n')
    p = tmp_path / "t.md"
    p.write_text("x\n")
    with pytest.raises(corecteaza.config.EroareUtilizare, match="systemctl --user start unelte-languagetool"):
        corecteaza.main([str(p)])


# ---------- traduce ----------

DOC = """---
titlu: Articol
---

# Lichiditatea

Formula $L_c = AC / DC$ arată **lichiditatea** [@pop2020].

| Indicator | Valoare |
|---|---|
| Lc | 2 |

```python
x = "nu se traduce"
```

$$
L_r = \\frac{AC - S}{DC}
$$

- primul punct
- al doilea `cod`
"""


def fake_translator(fake_ai, transform=lambda t: f"EN:{t}", drop=False):
    """Răspunde la fiecare cerere cu blocurile primite, „traduse”."""
    def respond(body):
        sent = body["messages"][-1]["content"]
        blocks = json.loads(sent.split("(JSON):\n", 1)[1])["blocuri"]
        out = [{"id": b["id"], "text": transform(b["text"])} for b in blocks]
        if drop:
            out = out[:-1]
        return json.dumps({"blocuri": out}, ensure_ascii=False)
    H = fake_ai.httpd.RequestHandlerClass

    def do_post(self):
        n = int(self.headers.get("Content-Length", 0))
        body = json.loads(self.rfile.read(n))
        fake_ai.requests.append({"method": "POST", "json": body})
        self._send(200, fake_ai.text_body(respond(body)))
    H.do_POST = do_post


def test_cod_si_formule_identice(env, fake_ai, tmp_path):
    fake_translator(fake_ai)
    p = tmp_path / "a.md"
    p.write_text(DOC)
    assert traduce.main([str(p), "--in", "en"]) == 0
    out = (tmp_path / "a.en.md").read_text()
    assert out.startswith("---\ntitlu: Articol\n---\n")
    assert '```python\nx = "nu se traduce"\n```' in out
    assert "$$\nL_r = \\frac{AC - S}{DC}\n$$" in out
    assert "EN:Formula $L_c = AC / DC$ arată **lichiditatea** [@pop2020]." in out
    assert "EN:- primul punct\n- al doilea `cod`" in out
    sent = fake_ai.chat_requests[0]["json"]["messages"][-1]["content"]
    assert "L_c" not in sent and "nu se traduce" not in sent and "⟦0⟧" in sent
    assert p.read_text() == DOC


def test_nepotrivire_detectata_fara_fisier(env, fake_ai, tmp_path):
    fake_translator(fake_ai, drop=True)
    p = tmp_path / "a.md"
    p.write_text("Unu.\n\nDoi.\n")
    with pytest.raises(traduce.config.EroareUtilizare, match="nu corespunde"):
        traduce.main([str(p), "--in", "en"])
    assert len(fake_ai.chat_requests) == 2                   # retrimisă o dată
    assert not (tmp_path / "a.en.md").exists()


def test_formula_pierduta_e_nepotrivire(env, fake_ai, tmp_path):
    fake_translator(fake_ai, transform=lambda t: "fără marcaje")
    with pytest.raises(traduce.config.EroareUtilizare):
        traduce.translate_markdown("Formula $x$ aici.\n", "en", None, None)


def test_glosar_in_prompt(env, fake_ai, tmp_path):
    (env / "glosar.toml").write_text('[en]\n"active circulante" = "current assets"\n')
    fake_translator(fake_ai)
    traduce.translate_markdown("Activele circulante.\n", "en", "ro", None)
    system = fake_ai.chat_requests[0]["json"]["messages"][0]["content"]
    assert "„active circulante” → „current assets”" in system and "în engleză" in system


def test_docx(env, fake_ai, tmp_path):
    from docx import Document
    fake_translator(fake_ai)
    d = Document()
    d.add_paragraph("Primul paragraf", style="Heading 1")
    d.add_paragraph("Al doilea")
    t = d.add_table(rows=1, cols=1)
    t.cell(0, 0).text = "Celulă"
    src = tmp_path / "d.docx"
    d.save(str(src))
    assert traduce.main([str(src), "--in", "en"]) == 0
    out = Document(str(tmp_path / "d.en.docx"))
    assert out.paragraphs[0].text == "EN:Primul paragraf" and out.paragraphs[0].style.name == "Heading 1"
    assert out.tables[0].cell(0, 0).text == "EN:Celulă"
