"""Generează reference.docx pentru un profil, coerent cu profil.toml.

Pornește de la șablonul implicit pandoc și aplică fonturile, mărimile, spațierile,
marginile și numerotarea paginilor din profil. Rulează după ce modifici profil.toml:

    .venv/bin/python profiluri/genereaza_reference.py profiluri/facultate

Poți și să deschizi reference.docx în Word/LibreOffice și să-i modifici stilurile manual;
`unelte doctor` verifică apoi coerența cu profil.toml.
"""
from __future__ import annotations

import subprocess
import sys
import tempfile
import tomllib
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ALIGN = {"justify": WD_ALIGN_PARAGRAPH.JUSTIFY, "left": WD_ALIGN_PARAGRAPH.LEFT,
         "center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT}
PAGE_SIZES = {"A4": (21.0, 29.7), "A5": (14.8, 21.0), "Letter": (21.59, 27.94)}
BODY_STYLES = ("Body Text", "First Paragraph")


def set_fonts(rpr_parent, font: str) -> None:
    """Fontul explicit pe toate scripturile, fără fonturi de temă (care ar avea prioritate)."""
    rpr = rpr_parent.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if fonts.get(qn(attr)) is not None:
            del fonts.attrib[qn(attr)]
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), font)


def set_doc_defaults(doc, font: str, size_pt: float) -> None:
    styles = doc.styles.element
    defaults = styles.find(qn("w:docDefaults"))
    if defaults is None:
        defaults = OxmlElement("w:docDefaults")
        styles.insert(0, defaults)
    rpr_default = defaults.find(qn("w:rPrDefault"))
    if rpr_default is None:
        rpr_default = OxmlElement("w:rPrDefault")
        defaults.insert(0, rpr_default)

    class Holder:
        def get_or_add_rPr(self):
            rpr = rpr_default.find(qn("w:rPr"))
            if rpr is None:
                rpr = OxmlElement("w:rPr")
                rpr_default.append(rpr)
            return rpr

    h = Holder()
    set_fonts(h, font)
    rpr = h.get_or_add_rPr()
    for tag in ("w:sz", "w:szCs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            rpr.append(el)
        el.set(qn("w:val"), str(int(round(size_pt * 2))))


def style(doc, name: str, base: str = "Normal"):
    for st in doc.styles:
        if st.name == name:
            return st
    else:
        st = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        st.base_style = doc.styles[base]
        st.quick_style = True
        return st


def add_page_number(section) -> None:
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    for fld in p._p.findall(qn("w:fldSimple")):
        p._p.remove(fld)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "1"
    run.append(t)
    fld.append(run)
    p._p.append(fld)


def build(profile_dir: Path) -> Path:
    with (profile_dir / "profil.toml").open("rb") as f:
        prof = tomllib.load(f)
    text, heads, page = prof.get("text", {}), prof.get("titluri", {}), prof.get("pagina", {})
    font = text.get("font", "Times New Roman")
    size = float(text.get("marime_pt", 12))
    spacing = float(text.get("spatiere_rand", 1.5))
    align = ALIGN.get(text.get("aliniere", "justify"), WD_ALIGN_PARAGRAPH.JUSTIFY)
    indent = float(text.get("indent_prima_linie_cm", 1.25))

    with tempfile.TemporaryDirectory() as tmp:
        base = Path(tmp) / "base.docx"
        with base.open("wb") as f:
            subprocess.run(["pandoc", "--print-default-data-file", "reference.docx"], stdout=f, check=True)
        doc = Document(str(base))

    set_doc_defaults(doc, font, size)
    normal = next(x for x in doc.styles if x.name == "Normal")
    set_fonts(normal.element, font)
    normal.font.size = Pt(size)
    normal.paragraph_format.line_spacing = spacing

    for name in BODY_STYLES:
        st = style(doc, name)
        set_fonts(st.element, font)
        st.font.size = Pt(size)
        pf = st.paragraph_format
        pf.alignment = align
        pf.first_line_indent = Cm(indent)
        pf.line_spacing = spacing
        pf.space_before = Pt(0)
        pf.space_after = Pt(6)

    hfont = heads.get("font", font)
    sizes = heads.get("marimi_pt", [16, 14, 12])
    for level in range(1, 10):
        st = next((x for x in doc.styles if x.name and x.name.lower() == f"heading {level}"), None)
        if st is None:
            continue
        set_fonts(st.element, hfont)
        st.font.size = Pt(float(sizes[min(level, len(sizes)) - 1]))
        st.font.bold = True
        st.font.italic = False
        st.font.color.rgb = RGBColor(0, 0, 0)
        pf = st.paragraph_format
        pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf.first_line_indent = Cm(0)
        pf.space_before = Pt(12 if level == 1 else 10)
        pf.space_after = Pt(6)
        pf.keep_with_next = True
        pf.line_spacing = 1.15

    for st in doc.styles:
        if st.name in ("Title", "Subtitle", "Author", "Date", "Abstract", "TOC Heading"):
            set_fonts(st.element, hfont)
            st.font.color.rgb = RGBColor(0, 0, 0)

    tp = style(doc, "Pagina titlu")
    set_fonts(tp.element, font)
    tp.font.size = Pt(14)
    tp.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tp.paragraph_format.first_line_indent = Cm(0)
    tp.paragraph_format.space_after = Pt(0)
    tl = style(doc, "Titlu lucrare", "Pagina titlu")
    set_fonts(tl.element, hfont)
    tl.font.size = Pt(20)
    tl.font.bold = True
    tl.paragraph_format.space_before = Pt(12)
    tl.paragraph_format.space_after = Pt(12)

    for name in ("Caption", "Image Caption", "Table Caption"):
        st = style(doc, name)
        set_fonts(st.element, font)
        st.font.size = Pt(max(size - 1, 9))
        st.font.italic = True
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        st.paragraph_format.first_line_indent = Cm(0)
    for name in ("Captioned Figure", "Figure"):
        st = style(doc, name)
        st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        st.paragraph_format.first_line_indent = Cm(0)
    bib = style(doc, "Bibliography")
    set_fonts(bib.element, font)
    bib.font.size = Pt(size)
    bib.paragraph_format.first_line_indent = Cm(-1.0)
    bib.paragraph_format.left_indent = Cm(1.0)
    bib.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    w, h = PAGE_SIZES.get(page.get("format", "A4"), PAGE_SIZES["A4"])
    margins = page.get("margini_cm", {})
    for section in doc.sections:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width, section.page_height = Cm(w), Cm(h)
        section.top_margin = Cm(float(margins.get("sus", 2.5)))
        section.bottom_margin = Cm(float(margins.get("jos", 2.5)))
        section.left_margin = Cm(float(margins.get("stanga", 3.0)))
        section.right_margin = Cm(float(margins.get("dreapta", 2.0)))
        if page.get("numerotare", True):
            add_page_number(section)

    out = profile_dir / "reference.docx"
    doc.save(str(out))
    return out


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit("Utilizare: genereaza_reference.py <folderul profilului>")
    print(f"Scris: {build(Path(sys.argv[1]))}")
